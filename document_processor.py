from pathlib import Path
import hashlib
import io
import re

import pymupdf as fitz
from docx import Document


ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}


def document_id_from_bytes(data):
    return hashlib.sha256(data).hexdigest()


def _clean(text):
    lines = [line.strip() for line in text.replace("\x00", "").splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _pdf_pages(data):
    pages = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = _clean(page.get_text("text"))
            if text:
                pages.append({"text": text, "page": page_number})
    return pages


def _docx_text(data):
    doc = Document(io.BytesIO(data))
    paragraphs = [_clean(p.text) for p in doc.paragraphs if _clean(p.text)]
    text = "\n".join(paragraphs)

    # Tables are also useful document content.
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append(" | ".join(_clean(cell.text) for cell in row.cells))
        table_text = "\n".join(r for r in rows if r.strip())
        if table_text:
            text += ("\n" if text else "") + table_text

    return [{"text": _clean(text), "page": None}] if text.strip() else []


def _txt_text(data):
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("utf-8", errors="replace")
    text = _clean(text)
    return [{"text": text, "page": None}] if text else []


def _split_text(text, chunk_size=800, chunk_overlap=100):
    """Splits text into chunks respecting sentence boundaries without external dependencies."""
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    sentences = re.split(r'(?<=[.!?\n])\s+', text)
    chunks = []
    current_chunk = []
    current_len = 0

    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue

        if len(sentence) > chunk_size:
            words = sentence.split()
            sub_chunk = []
            sub_len = 0
            for word in words:
                if sub_len + len(word) + 1 > chunk_size and sub_chunk:
                    chunks.append(" ".join(sub_chunk))
                    overlap_words = []
                    overlap_len = 0
                    for w in reversed(sub_chunk):
                        if overlap_len + len(w) + 1 <= chunk_overlap:
                            overlap_words.insert(0, w)
                            overlap_len += len(w) + 1
                        else:
                            break
                    sub_chunk = overlap_words + [word]
                    sub_len = sum(len(w) for w in sub_chunk) + len(sub_chunk) - 1
                else:
                    sub_chunk.append(word)
                    sub_len += len(word) + 1
            if sub_chunk:
                chunks.append(" ".join(sub_chunk))
            continue

        if current_len + len(sentence) + 1 > chunk_size and current_chunk:
            chunks.append(" ".join(current_chunk))

            overlap_sentences = []
            overlap_len = 0
            for s in reversed(current_chunk):
                if overlap_len + len(s) + 1 <= chunk_overlap:
                    overlap_sentences.insert(0, s)
                    overlap_len += len(s) + 1
                else:
                    break
            current_chunk = overlap_sentences + [sentence]
            current_len = sum(len(s) for s in current_chunk) + len(current_chunk) - 1
        else:
            current_chunk.append(sentence)
            current_len += len(sentence) + 1

    if current_chunk:
        chunks.append(" ".join(current_chunk))

    return chunks


def process_document(data, filename, chunk_size=800, chunk_overlap=100):
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Only PDF, DOCX, and TXT files are allowed.")

    document_id = document_id_from_bytes(data)

    if ext == "pdf":
        page_parts = _pdf_pages(data)
    elif ext == "docx":
        page_parts = _docx_text(data)
    else:
        page_parts = _txt_text(data)

    if not page_parts:
        raise ValueError("The uploaded document does not contain readable text.")

    chunks = []
    chunk_index = 0

    for part in page_parts:
        nodes = _split_text(part["text"], chunk_size=chunk_size, chunk_overlap=chunk_overlap)
        for node_text in nodes:
            cleaned = _clean(node_text)
            if not cleaned:
                continue
            chunks.append({
                "text": cleaned,
                "source": filename,
                "file_type": ext,
                "page": part["page"],
                "chunk_index": chunk_index,
                "document_id": document_id,
            })
            chunk_index += 1

    if not chunks:
        raise ValueError("The uploaded document does not contain readable text.")

    return {
        "document_id": document_id,
        "filename": filename,
        "file_type": ext,
        "chunks": chunks,
    }
